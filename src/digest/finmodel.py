"""Per-security operating model: a simplified, driver-based version of a
Canalyst model (quarters + fiscal years, non-GAAP).

Two halves:
  * parse_canalyst(path)  -> a *base* document: periods, segment revenue, P&L
    lines, cash flow and balance items, actuals and the vendor's forecasts.
  * compute(base, overrides) -> the rendered model. Every forecast period is
    rebuilt from drivers (segment y/y growth, gross margin, R&D and SG&A as % of
    revenue, net other income, tax rate, diluted shares, OCF conversion, capex).
    Drivers default to the values implied by the base, so with no overrides the
    output reproduces the vendor model; a user's overrides flow through the P&L
    to EPS, FCF and implied value.

Bases are shared per company (model_bases table); overrides are per user
(security_models table, see model.py).
"""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path

_PERIOD_RE = re.compile(r"^(Q([1-4])-(\d{4})|FY(\d{4}))$")
_KEEP_ACTUAL_QUARTERS = 8      # history kept in the base (q/q and y/y need it)
_SHOW_ACTUAL_QUARTERS = 4      # history shown in the rendered model
_SHOW_FORECAST_QUARTERS = 6
_SHOW_ACTUAL_YEARS = 1
_SHOW_FORECAST_YEARS = 3

# Canalyst row labels for the company-level lines (first match wins).
_LINES: dict[str, tuple[str, ...]] = {
    "revenue": ("Total Revenue, mm", "Revenue, mm", "Net Revenue, mm"),
    "gross_profit": ("Non-GAAP Gross Profit, mm", "GAAP Gross Profit, mm", "Gross Profit, mm"),
    "rd": ("Non-GAAP R&D Expense, mm", "GAAP R&D Expense, mm", "R&D Expense, mm"),
    "sga": ("Non-GAAP SG&A Expense, mm", "GAAP SG&A Expense, mm", "SG&A Expense, mm"),
    "operating_income": ("Non-GAAP Operating Income, mm", "Company Reported Non-GAAP Operating Income",
                         "GAAP Operating Income, mm", "Operating Income, mm"),
    "ebt": ("Non-GAAP Earnings before Taxes", "GAAP Earnings before Taxes"),
    "tax": ("Non-GAAP Provision for taxes", "GAAP Provision for Taxes"),
    "net_income": ("Non-GAAP Net Income Attributable to Common Shareholders",
                   "GAAP Net Income - Diluted"),
    "shares": ("Adjusted Shares Outstanding - WAD", "Shares Outstanding - WAD"),
    "eps": ("Non-GAAP Adjusted Earnings Per Share - WAD", "GAAP Earnings Per Share - WAD"),
    "da": ("Depreciation & Amortization, mm",),
    "ocf": ("Operating Cash Flow after WC",),
    "capex": ("Capex", "Capex - reported/model estimate, mm"),
}


# ── parsing ────────────────────────────────────────────────────────────────


def _num(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def parse_canalyst(path: str | Path) -> dict:
    """Read a Canalyst-format workbook into a base model document."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Model" not in wb.sheetnames:
        raise ValueError("not a Canalyst-format model (no 'Model' sheet)")
    rows = list(wb["Model"].iter_rows(values_only=True))
    hdr_i = next(
        (i for i, r in enumerate(rows[:15])
         if sum(1 for c in r if isinstance(c, str) and _PERIOD_RE.match(c)) > 8),
        None,
    )
    if hdr_i is None:
        raise ValueError("not a Canalyst-format model (no Q1-2027 / FY2027 period header row)")
    header = rows[hdr_i]
    ends = next(
        (r for r in rows[:hdr_i] if sum(isinstance(c, datetime) for c in r) > 8), [None] * len(header)
    )
    cols: list[dict] = []
    for ci, h in enumerate(header):
        m = _PERIOD_RE.match(h) if isinstance(h, str) else None
        if not m:
            continue
        end = ends[ci].date() if ci < len(ends) and isinstance(ends[ci], datetime) else None
        if m.group(2):
            cols.append({"ci": ci, "id": h, "kind": "Q", "q": int(m.group(2)), "fy": int(m.group(3)), "end": end})
        else:
            cols.append({"ci": ci, "id": h, "kind": "FY", "fy": int(m.group(4)), "end": end})

    labelled = [(i, r[0].strip(), r) for i, r in enumerate(rows) if isinstance(r[0], str)]

    # Actual vs forecast: reported periods carry far more populated rows (the
    # as-reported/supplemental sections are historical-only).
    density = {c["id"]: sum(1 for _, _, r in labelled if _num(r[c["ci"]]) is not None) for c in cols}
    q_cols = [c for c in cols if c["kind"] == "Q"]
    peak = max(density[c["id"]] for c in q_cols)
    actual_q = [c for c in q_cols if density[c["id"]] >= 0.8 * peak]
    last_actual = actual_q[-1]
    last_key = (last_actual["fy"], last_actual["q"])

    def is_actual(c: dict) -> bool:
        if c["kind"] == "Q":
            return (c["fy"], c["q"]) <= last_key
        return c["fy"] < last_actual["fy"] or (c["fy"] == last_actual["fy"] and last_actual["q"] == 4)

    # Keep recent history plus every forecast period.
    keep_q = [c for c in q_cols if not is_actual(c)]
    keep_q = [c for c in actual_q[-_KEEP_ACTUAL_QUARTERS:]] + keep_q
    first_fy = keep_q[0]["fy"] - 1
    keep_fy = [c for c in cols if c["kind"] == "FY" and c["fy"] >= first_fy]
    kept = sorted(keep_q + keep_fy, key=lambda c: (c["fy"], c.get("q", 5)))

    def series(r) -> dict[str, float]:
        out = {}
        for c in kept:
            v = _num(r[c["ci"]]) if c["ci"] < len(r) else None
            if v is not None:
                out[c["id"]] = v
        return out

    lines: dict[str, dict[str, float]] = {}
    for key, candidates in _LINES.items():
        for cand in candidates:
            hit = next((r for _, lab, r in labelled if lab.lower() == cand.lower()), None)
            if hit is not None and series(hit):
                lines[key] = series(hit)
                break
    if "capex" in lines:
        lines["capex"] = {k: abs(v) for k, v in lines["capex"].items()}

    # Segment revenue: the revenue-build rows before the first total that the
    # vendor actually forecasts.
    first_fc = next((c["id"] for c in kept if not is_actual(c)), None)
    total_i = next((i for i, lab, _ in labelled if lab.lower() == "total revenue, mm"), None)
    if first_fc is None or total_i is None:
        raise ValueError("not a Canalyst-format model (no forecast periods or 'Total Revenue, mm' row)")
    segs: list[dict] = []
    for i, lab, r in labelled:
        if i >= total_i:
            break
        if not lab.lower().endswith(", mm") or lab.lower().startswith(("consensus", "guidance", "total")):
            continue
        s = series(r)
        if first_fc not in s:
            continue
        name = re.sub(r"(?i)\s*(revenue)?,\s*mm$", "", lab).strip()
        segs.append({"id": f"s{len(segs)}", "name": name, "revenue": s})
    _mark_subtotals(segs, kept)
    _add_residual_segment(segs, lines.get("revenue", {}))

    balance = _balance(wb, last_actual["id"])
    company = rows[0][0] if isinstance(rows[0][0], str) else None
    ticker = rows[1][0].split()[0] if isinstance(rows[1][0], str) else None
    return {
        "source": "Canalyst",
        "file": Path(path).name,
        "company_name": company,
        "ticker": ticker,
        "periods": [
            {k: (v.isoformat() if isinstance(v, date) else v) for k, v in c.items() if k != "ci"}
            | {"actual": is_actual(c)}
            for c in kept
        ],
        "segments": segs,
        "lines": lines,
        "balance": balance,
        "imported_at": datetime.utcnow().isoformat(),
    }


def _mark_subtotals(segs: list[dict], periods: list[dict]) -> None:
    """A segment equal to the sum of the items just before it is a subtotal.
    Nested subtotals work: a subtotal can sum earlier subtotals."""
    ids = [p["id"] for p in periods]
    pending: list[dict] = []  # leaves and subtotals not yet rolled up
    for s in segs:
        for n in range(len(pending), 1, -1):
            group = pending[-n:]
            checks = [
                (sum(g["revenue"].get(pid, 0.0) for g in group), s["revenue"][pid])
                for pid in ids if pid in s["revenue"]
            ]
            if checks and any(abs(b) > 0 for _, b in checks) and all(
                abs(a - b) <= 0.005 * max(1.0, abs(b)) for a, b in checks
            ):
                s["children"] = [g["id"] for g in group]
                pending = pending[:-n]
                break
        pending.append(s)


# ── ConvexityAI template ("Model" + "Consensus Model" sheets) ────────────────

_CVX_LINES: dict[str, tuple[str, ...]] = {
    # line -> label prefixes (case-insensitive), first match wins
    "revenue": ("net revenue", "total revenue", "revenue"),
    "gross_profit": ("gross profit",),
    "sga": ("selling, general", "sg&a"),
    "rd": ("research & development", "r&d"),
    "operating_income": ("ebit", "operating income", "operating profit"),
    "da": ("add back: depreciation", "depreciation & amortization", "d&a"),
    "ebt": ("pre-tax income", "income before tax", "ebt"),
    "tax": ("income tax provision", "provision for income taxes", "tax expense"),
    "net_income": ("net income for eps", "net income attributable", "net income"),
    "shares": ("diluted shares",),
    "eps": ("diluted eps", "eps"),
    "ocf": ("operating cash flow", "cash from operations"),
    "capex": ("capex", "capital expenditure"),
}


# Ratio rows ("Y/Y growth, %", "SG&A, % of revenue") — not "...(40% stake)".
_PCT_ROW = re.compile(r"growth|,\s*%|%\s*$|% of ")


def is_convexity_template(wb) -> bool:
    if "Model" not in wb.sheetnames:
        return False
    for r in wb["Model"].iter_rows(max_row=40, values_only=True):
        if any(isinstance(c, str) and c.strip().upper() == "REVENUE BY SEGMENT" for c in r):
            return True
    return False


def _parse_cvx_sheet(ws) -> dict | None:
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(
        (i for i, r in enumerate(rows[:15])
         if sum(1 for c in r if isinstance(c, str) and _PERIOD_RE.match(c.strip())) >= 4),
        None,
    )
    if hdr_i is None:
        return None
    header = rows[hdr_i]
    cols = []
    for ci, h in enumerate(header):
        m = _PERIOD_RE.match(h.strip()) if isinstance(h, str) else None
        if not m:
            continue
        if m.group(2):
            cols.append({"ci": ci, "id": h.strip(), "kind": "Q", "q": int(m.group(2)), "fy": int(m.group(3))})
        else:
            cols.append({"ci": ci, "id": h.strip(), "kind": "FY", "fy": int(m.group(4))})
    marker = next((r for r in rows[hdr_i + 1 : hdr_i + 3]
                   if any(isinstance(c, str) and c.strip().lower() == "forecast" for c in r)), None)
    fc_col = next((ci for ci, c in enumerate(marker) if isinstance(c, str) and c.strip().lower() == "forecast"),
                  None) if marker else None
    first_col = cols[0]["ci"]

    def label_of(r) -> str | None:
        for c in r[:first_col]:
            if isinstance(c, str) and c.strip():
                return c.strip()
        return None

    def series(r) -> dict[str, float]:
        return {c["id"]: v for c in cols if c["ci"] < len(r) and (v := _num(r[c["ci"]])) is not None}

    labelled = [(label_of(r), r) for r in rows[hdr_i + 1 :]]
    labelled = [(lab, r) for lab, r in labelled if lab]

    lines: dict[str, dict[str, float]] = {}
    for key, prefixes in _CVX_LINES.items():
        for lab, r in labelled:
            low = lab.lower()
            if any(low == p or low.startswith(p) for p in prefixes) and not _PCT_ROW.search(low) and "margin" not in low:
                ser = series(r)
                if ser:
                    lines[key] = ser
                    break
    if "capex" in lines:
        lines["capex"] = {k: abs(v) for k, v in lines["capex"].items()}

    # Segment revenue: "... revenue" rows between REVENUE BY SEGMENT and the total.
    segs: list[dict] = []
    in_seg = False
    for lab, r in labelled:
        low = lab.lower()
        if low == "revenue by segment":
            in_seg = True
            continue
        if not in_seg:
            continue
        if low.startswith(("net revenue", "total revenue")) or low in ("income statement",):
            break
        if "revenue" in low and not _PCT_ROW.search(low):
            ser = series(r)
            if ser:
                name = re.sub(r"(?i)\s*revenue\b", "", lab).strip(" ,")
                segs.append({"id": f"s{len(segs)}", "name": name or lab, "revenue": ser})

    return {"cols": cols, "fc_col": fc_col, "lines": lines, "segments": segs}


def parse_convexity(path: str | Path) -> dict:
    """Read a ConvexityAI-template workbook: 'Model' (house/user forecast) and an
    optional 'Consensus Model' sheet in the same layout."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    model = _parse_cvx_sheet(wb["Model"])
    if model is None or not model["lines"].get("revenue"):
        raise ValueError("ConvexityAI template: no period header row or revenue line found in 'Model'")
    fc_col = model["fc_col"]

    def is_actual(c: dict) -> bool:
        return fc_col is None or c["ci"] < fc_col

    # Everything left of the "Forecast" marker is actual — fiscal-year columns
    # included (a year whose last quarter is a forecast sits right of it).
    cols = model["cols"]
    periods = [
        {"id": c["id"], "kind": c["kind"], "fy": c["fy"], **({"q": c["q"]} if c["kind"] == "Q" else {}),
         "end": None, "actual": is_actual(c)}
        for c in cols
    ]
    lines = model["lines"]
    _fill_eps(lines)
    segs = model["segments"]
    _mark_subtotals(segs, periods)
    _add_residual_segment(segs, lines.get("revenue", {}))

    consensus = None
    if "Consensus Model" in wb.sheetnames:
        cons = _parse_cvx_sheet(wb["Consensus Model"])
        if cons and cons["lines"]:
            _fill_eps(cons["lines"])
            consensus = {"lines": cons["lines"],
                         "segments": {s["name"]: s["revenue"] for s in cons["segments"]}}

    title = next((c for r in wb["Model"].iter_rows(max_row=3, values_only=True) for c in r
                  if isinstance(c, str) and c.strip()), None)
    return {
        "source": "ConvexityAI model",
        "file": Path(path).name,
        "company_name": title,
        "ticker": None,
        "periods": periods,
        "segments": segs,
        "lines": lines,
        "balance": {},
        "consensus": consensus,
        "imported_at": datetime.utcnow().isoformat(),
    }


def _fill_eps(lines: dict) -> None:
    """EPS = net income / diluted shares wherever the sheet leaves EPS blank."""
    ni, sh = lines.get("net_income", {}), lines.get("shares", {})
    eps = lines.setdefault("eps", {})
    for pid, v in ni.items():
        if pid not in eps and sh.get(pid):
            eps[pid] = v / sh[pid]


def parse_model(path: str | Path) -> dict:
    """Any supported model workbook: the ConvexityAI template or a Canalyst model."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    return parse_convexity(path) if is_convexity_template(wb) else parse_canalyst(path)


def _add_residual_segment(segs: list[dict], total: dict[str, float]) -> None:
    leaves = [s for s in segs if "children" not in s]
    # Only periods where every segment reports (segment disclosure changes over time).
    resid = {
        pid: total[pid] - sum(s["revenue"][pid] for s in leaves)
        for pid in total
        if all(pid in s["revenue"] for s in leaves)
    }
    if any(abs(v) > 0.005 * max(1.0, abs(total[pid])) for pid, v in resid.items()):
        segs.append({"id": f"s{len(segs)}", "name": "Other", "revenue": resid})


def _balance(wb, period_id: str) -> dict:
    """Cash and debt at the last reported quarter (Summary Page capitalization block)."""
    if "Summary Page" not in wb.sheetnames:
        return {}
    rows = list(wb["Summary Page"].iter_rows(values_only=True))
    ci = next((r.index(period_id) for r in rows[:6] if period_id in r), None)
    if ci is None:
        return {}
    out: dict = {"as_of": period_id}
    for r in rows:
        lab = r[0].strip() if isinstance(r[0], str) else None
        if lab in ("Cash", "Debt") and lab.lower() not in out and ci < len(r):
            val = _num(r[ci])
            if val is not None:
                out[lab.lower()] = val
    return out


# ── compute ────────────────────────────────────────────────────────────────


def _div(a: float | None, b: float | None) -> float | None:
    return a / b if a is not None and b not in (None, 0) else None


def _prior(pid: str, by_id: dict[str, dict], kind: str, step: str) -> str | None:
    p = by_id[pid]
    if p["kind"] == "FY":
        key = f"FY{p['fy'] - 1}"
    elif step == "yoy":
        key = f"Q{p['q']}-{p['fy'] - 1}"
    else:
        key = f"Q{p['q'] - 1}-{p['fy']}" if p["q"] > 1 else f"Q4-{p['fy'] - 1}"
    return key if key in by_id else None


def _label(p: dict) -> tuple[str, str]:
    """('3Q27E', '3Q26') — fiscal label and calendar label (mid-quarter rule)."""
    e = "" if p["actual"] else "E"
    end = date.fromisoformat(p["end"]) if p.get("end") else None
    if p["kind"] == "FY":
        cal = f"CY{(end - timedelta(days=182)).year}" if end else ""
        return f"FY{p['fy']}{e}", cal
    if end:
        mid = end - timedelta(days=45)
        cal = f"{(mid.month - 1) // 3 + 1}Q{str(mid.year)[2:]}"
    else:
        cal = ""
    return f"{p['q']}Q{str(p['fy'])[2:]}{e}", cal


def compute(base: dict, overrides: dict[str, float] | None = None, price: float | None = None,
            target_pe: float | None = None) -> dict:
    """Rebuild every forecast period from drivers and render the model."""
    V, drivers, base_drivers = run_engine(base, overrides)
    return _render(base, V, drivers, base_drivers, overrides or {}, price, target_pe)


def run_engine(base: dict, overrides: dict[str, float] | None = None):
    """Recompute every period of *base* under *overrides*.

    Returns (V, drivers, base_drivers): V[key][period_id] -> value for keys
    'revenue', 'gross_profit', ..., 'seg:<id>'."""
    overrides = overrides or {}
    periods = base["periods"]
    by_id = {p["id"]: p for p in periods}
    L = base["lines"]
    segs = base["segments"]
    leaves = [s for s in segs if "children" not in s]

    quarters_by_fy: dict[int, list[str]] = {}
    for p in periods:
        if p["kind"] == "Q":
            quarters_by_fy.setdefault(p["fy"], []).append(p["id"])
    full_q_fy = {fy for fy, qs in quarters_by_fy.items() if len(qs) == 4}

    V: dict[str, dict[str, float | None]] = {}  # computed values: key -> pid -> value

    def v(key: str, pid: str | None) -> float | None:
        return V.get(key, {}).get(pid) if pid else None

    def put(key: str, pid: str, val: float | None) -> None:
        V.setdefault(key, {})[pid] = val

    drivers: dict[str, dict[str, float | None]] = {}  # driver key -> pid -> effective value
    base_drivers: dict[str, dict[str, float | None]] = {}

    def driver(key: str, pid: str, default: float | None) -> float | None:
        base_drivers.setdefault(key, {})[pid] = default
        val = overrides.get(f"{key}@{pid}", default)
        drivers.setdefault(key, {})[pid] = val
        return val

    def bval(key: str, pid: str) -> float | None:
        return L.get(key, {}).get(pid)

    order = sorted(periods, key=lambda p: (p["fy"], p.get("q", 5)))
    for p in order:
        pid = p["id"]
        derived_fy = p["kind"] == "FY" and p["fy"] in full_q_fy
        if p["actual"] or derived_fy:
            if derived_fy and not p["actual"]:
                _sum_fy(V, pid, quarters_by_fy[p["fy"]], leaves)
                continue
            for s in leaves:
                put(f"seg:{s['id']}", pid, s["revenue"].get(pid))
            for key in ("revenue", "gross_profit", "rd", "sga", "operating_income", "ebt", "tax",
                        "net_income", "shares", "eps", "da", "ocf", "capex"):
                put(key, pid, bval(key, pid))
            continue

        # ── forecast period: rebuild from drivers ──
        prev = _prior(pid, by_id, p["kind"], "yoy")
        rev = 0.0
        for s in leaves:
            k = f"seg:{s['id']}"
            prior_v = v(k, prev)
            # Default growth is the base model's own rate, so an upstream change
            # compounds forward instead of snapping back to the base level.
            default_g = _div(s["revenue"].get(pid), s["revenue"].get(prev) if prev else None)
            g = driver(f"growth:{s['id']}", pid, default_g - 1 if default_g is not None else None)
            val = prior_v * (1 + g) if prior_v is not None and g is not None else s["revenue"].get(pid)
            put(k, pid, val)
            rev += val or 0.0
        put("revenue", pid, rev)

        def pct(name: str, line: str) -> float | None:
            return driver(name, pid, _div(bval(line, pid), bval("revenue", pid)))

        gm = pct("gm", "gross_profit")
        rd_pct = pct("rd_pct", "rd")
        sga_pct = pct("sga_pct", "sga")
        gp = rev * gm if gm is not None else None
        rd = rev * rd_pct if rd_pct is not None else 0.0
        sga = rev * sga_pct if sga_pct is not None else 0.0
        b_other_opex = (bval("gross_profit", pid) or 0) - (bval("rd", pid) or 0) - (bval("sga", pid) or 0) \
            - (bval("operating_income", pid) or 0)
        oi = (gp or 0) - rd - sga - b_other_opex
        other = driver("other_income", pid, (bval("ebt", pid) or 0) - (bval("operating_income", pid) or 0))
        ebt = oi + (other or 0)
        tax_rate = driver("tax_rate", pid, _div(bval("tax", pid), bval("ebt", pid)))
        tax = ebt * tax_rate if tax_rate is not None else 0.0
        ni = ebt - tax
        shares = driver("shares", pid, bval("shares", pid))
        conv = driver("ocf_conv", pid, _div(bval("ocf", pid), bval("net_income", pid)))
        capex = driver("capex", pid, bval("capex", pid))
        for key, val in (("gross_profit", gp), ("rd", rd), ("sga", sga), ("operating_income", oi),
                         ("ebt", ebt), ("tax", tax), ("net_income", ni), ("shares", shares),
                         ("eps", _div(ni, shares)), ("da", bval("da", pid)),
                         ("ocf", ni * conv if conv is not None else None), ("capex", capex)):
            put(key, pid, val)

    for p in order:  # subtotal segments (children precede their subtotal)
        for s in segs:
            if "children" in s:
                put(f"seg:{s['id']}", p["id"],
                    sum(v(f"seg:{c}", p["id"]) or 0.0 for c in s["children"]) if p["id"] in V.get("revenue", {}) else None)

    return V, drivers, base_drivers


def _sum_fy(V: dict, fy_id: str, qids: list[str], leaves: list[dict]) -> None:
    def tot(key: str) -> float | None:
        vals = [V.get(key, {}).get(q) for q in qids]
        return sum(vals) if all(x is not None for x in vals) else None

    for key in ["revenue", "gross_profit", "rd", "sga", "operating_income", "ebt", "tax", "net_income",
                "da", "ocf", "capex"] + [f"seg:{s['id']}" for s in leaves]:
        V.setdefault(key, {})[fy_id] = tot(key)
    sh = [V.get("shares", {}).get(q) for q in qids]
    V.setdefault("shares", {})[fy_id] = sum(sh) / 4 if all(sh) else None
    V.setdefault("eps", {})[fy_id] = _div(V["net_income"][fy_id], V["shares"][fy_id])


def _render(base, V, drivers, base_drivers, overrides, price, target_pe) -> dict:
    periods = base["periods"]
    by_id = {p["id"]: p for p in periods}
    qs = [p for p in periods if p["kind"] == "Q"]
    fys = [p for p in periods if p["kind"] == "FY"]
    act_q = [p for p in qs if p["actual"]]
    fc_q = [p for p in qs if not p["actual"]]
    act_fy = [p for p in fys if p["actual"]]
    fc_fy = [p for p in fys if not p["actual"]]
    shown = act_q[-_SHOW_ACTUAL_QUARTERS:] + fc_q[:_SHOW_FORECAST_QUARTERS] \
        + act_fy[-_SHOW_ACTUAL_YEARS:] + fc_fy[:_SHOW_FORECAST_YEARS]
    columns = []
    for p in shown:
        fiscal, cal = _label(p)
        columns.append({"id": p["id"], "label": fiscal, "calendar": cal, "kind": p["kind"],
                        "actual": p["actual"]})
    col_ids = [c["id"] for c in columns]

    def vals(key: str) -> dict:
        return {pid: V.get(key, {}).get(pid) for pid in col_ids}

    def ratio(a: str, b: str) -> dict:
        return {pid: _div(V.get(a, {}).get(pid), V.get(b, {}).get(pid)) for pid in col_ids}

    def growth(key: str, step: str) -> dict:
        out = {}
        for pid in col_ids:
            if step == "qoq" and by_id[pid]["kind"] == "FY":
                out[pid] = None
                continue
            prev = _prior(pid, by_id, by_id[pid]["kind"], step)
            cur, old = V.get(key, {}).get(pid), V.get(key, {}).get(prev) if prev else None
            out[pid] = cur / old - 1 if cur is not None and old not in (None, 0) else None
        return out

    def row(rid, label, values, fmt="num", style=None, driver=None):
        r = {"id": rid, "label": label, "fmt": fmt, "values": values}
        if style:
            r["style"] = style
        if driver:
            r["driver"] = driver
            r["base"] = {pid: base_drivers.get(driver, {}).get(pid) for pid in col_ids}
            r["overridden"] = [pid for pid in col_ids if f"{driver}@{pid}" in overrides]
        return r

    def driven(key: str, fallback: dict) -> dict:
        """Driver value in forecast quarters / annual-only years, computed elsewhere."""
        return {pid: drivers.get(key, {}).get(pid, fallback.get(pid)) if pid in drivers.get(key, {}) else fallback.get(pid)
                for pid in col_ids}

    summary = [
        row("revenue", "Revenues", vals("revenue"), style="bold"),
        row("revenue_qoq", "qoq", growth("revenue", "qoq"), "pct", "sub"),
        row("revenue_yoy", "yoy", growth("revenue", "yoy"), "pct", "sub"),
        row("gross_profit", "Gross Profit", vals("gross_profit"), style="bold"),
        row("gpm", "GPM %", ratio("gross_profit", "revenue"), "pct", "margin"),
        row("operating_income", "Operating Profit", vals("operating_income"), style="bold"),
        row("opm", "OPM %", ratio("operating_income", "revenue"), "pct", "margin"),
        row("net_income", "Net Income", vals("net_income"), style="bold"),
        row("npm", "NPM %", ratio("net_income", "revenue"), "pct", "margin"),
        row("eps", "EPS ($)", vals("eps"), "eps", "bold"),
    ]

    build: list[dict] = []
    for s in base["segments"]:
        k = f"seg:{s['id']}"
        is_sub = "children" in s
        build.append(row(k, f"{s['name']} revenue", vals(k), style="bold" if not is_sub else "total"))
        build.append(row(f"{k}:qoq", "qoq", growth(k, "qoq"), "pct", "sub"))
        if is_sub:
            build.append(row(f"{k}:yoy", "yoy", growth(k, "yoy"), "pct", "sub"))
        else:
            build.append(row(f"{k}:yoy", "yoy", driven(f"growth:{s['id']}", growth(k, "yoy")), "pct", "sub",
                             driver=f"growth:{s['id']}"))
    build.append(row("revenue_total", "Total revenue", vals("revenue"), style="total"))

    pnl = [
        row("p_revenue", "Revenue", vals("revenue"), style="bold"),
        row("p_gp", "Gross profit", vals("gross_profit")),
        row("p_gm", "Gross margin %", driven("gm", ratio("gross_profit", "revenue")), "pct", "margin", driver="gm"),
        row("p_rd", "R&D", vals("rd")),
        row("p_rd_pct", "% of revenue", driven("rd_pct", ratio("rd", "revenue")), "pct", "sub", driver="rd_pct"),
        row("p_sga", "SG&A", vals("sga")),
        row("p_sga_pct", "% of revenue", driven("sga_pct", ratio("sga", "revenue")), "pct", "sub", driver="sga_pct"),
        row("p_oi", "Operating income", vals("operating_income"), style="bold"),
        row("p_opm", "Operating margin %", ratio("operating_income", "revenue"), "pct", "margin"),
        row("p_other", "Net interest & other", {pid: _sub(V, "ebt", "operating_income", pid) for pid in col_ids},
            driver="other_income"),
        row("p_ebt", "Pre-tax income", vals("ebt")),
        row("p_tax_rate", "Tax rate %", driven("tax_rate", ratio("tax", "ebt")), "pct", "sub", driver="tax_rate"),
        row("p_ni", "Net income", vals("net_income"), style="bold"),
        row("p_shares", "Diluted shares (mm)", driven("shares", vals("shares")), driver="shares"),
        row("p_eps", "EPS ($)", vals("eps"), "eps", "bold"),
        row("p_eps_yoy", "yoy", growth("eps", "yoy"), "pct", "sub"),
    ]

    fcf = {pid: _sub(V, "ocf", "capex", pid) for pid in col_ids}
    cash = [
        row("c_ebitda", "EBITDA", {pid: _add(V, "operating_income", "da", pid) for pid in col_ids}),
        row("c_ocf", "Operating cash flow", vals("ocf")),
        row("c_conv", "OCF / net income", driven("ocf_conv", ratio("ocf", "net_income")), "pct", "sub", driver="ocf_conv"),
        row("c_capex", "Capex", driven("capex", vals("capex")), driver="capex"),
        row("c_fcf", "Free cash flow", fcf, style="bold"),
        row("c_fcf_margin", "FCF margin %", {pid: _div(fcf[pid], V.get("revenue", {}).get(pid)) for pid in col_ids},
            "pct", "margin"),
        row("c_fcf_ps", "FCF / share ($)", {pid: _div(fcf[pid], V.get("shares", {}).get(pid)) for pid in col_ids}, "eps"),
    ]

    # Valuation on fiscal-year columns.
    bal = base.get("balance", {})
    shares_now = next((V["shares"][p["id"]] for p in reversed(qs) if p["actual"] and V.get("shares", {}).get(p["id"])), None)
    mcap = price * shares_now if price and shares_now else None
    net_cash = (bal.get("cash") or 0) - (bal.get("debt") or 0) if bal else 0
    ev = mcap - net_cash if mcap is not None else None
    fy_cols = [c["id"] for c in columns if c["kind"] == "FY"]
    ebitda = {pid: _add(V, "operating_income", "da", pid) for pid in fy_cols}
    pe = {pid: _div(price, V.get("eps", {}).get(pid)) for pid in fy_cols}
    implied = {pid: (V["eps"][pid] * target_pe) if target_pe and V.get("eps", {}).get(pid) else None for pid in fy_cols}
    valuation = [
        row("v_pe", "P/E", pe, "x"),
        row("v_ev_ebitda", "EV / EBITDA", {pid: _div(ev, ebitda[pid]) for pid in fy_cols}, "x"),
        row("v_ev_sales", "EV / Sales", {pid: _div(ev, V.get("revenue", {}).get(pid)) for pid in fy_cols}, "x"),
        row("v_fcf_yield", "FCF yield %", {pid: _div(fcf.get(pid), mcap) for pid in fy_cols}, "pct"),
        row("v_target_pe", "Target P/E", {pid: target_pe for pid in fy_cols if not by_id[pid]["actual"]}, "x", "bold"),
        row("v_implied", "Implied value ($)", implied, "eps", "bold"),
        row("v_upside", "Upside %", {pid: (implied[pid] / price - 1) if implied.get(pid) and price else None
                                     for pid in fy_cols}, "pct", "margin"),
    ]

    sections = [
        {"id": "summary", "title": "Simplified financials (non-GAAP, $mm except per share)", "rows": summary},
    ]
    cons = base.get("consensus")
    if cons and cons.get("lines"):
        CL = cons["lines"]

        def cval(line: str) -> dict:
            return {pid: CL.get(line, {}).get(pid) for pid in col_ids}

        def cratio(a: str, b: str) -> dict:
            return {pid: _div(CL.get(a, {}).get(pid), CL.get(b, {}).get(pid)) for pid in col_ids}

        def vs(line: str) -> dict:
            out = {}
            for pid in col_ids:
                mine, theirs = V.get(line, {}).get(pid), CL.get(line, {}).get(pid)
                out[pid] = (mine / theirs - 1) if mine is not None and theirs not in (None, 0) and not by_id[pid]["actual"] else None
            return out

        sections.append({"id": "consensus", "title": "Consensus vs your model", "rows": [
            row("c_rev", "Revenue — consensus", cval("revenue"), style="bold"),
            row("c_rev_vs", "Your revenue vs consensus", vs("revenue"), "pct", "variance"),
            row("c_gpm", "GPM % — consensus", cratio("gross_profit", "revenue"), "pct", "margin"),
            row("c_oi", "Operating profit — consensus", cval("operating_income")),
            row("c_oi_vs", "Your op. profit vs consensus", vs("operating_income"), "pct", "variance"),
            row("c_eps", "EPS — consensus", cval("eps"), "eps", "bold"),
            row("c_eps_vs", "Your EPS vs consensus", vs("eps"), "pct", "variance"),
        ]})
    sections += [
        {"id": "build", "title": "Revenue build", "rows": build},
        {"id": "pnl", "title": "P&L", "rows": pnl},
        {"id": "cash", "title": "Cash flow", "rows": cash},
        {"id": "valuation", "title": "Valuation", "rows": valuation, "annual_only": True},
    ]

    forecast_fy = [c["id"] for c in columns if c["kind"] == "FY" and not c["actual"]]
    line_rows = {"revenue": "revenue", "gross_margin": "gpm", "op_margin": "opm", "eps": "eps",
                 "multiple": "v_target_pe"}
    impact_cells = {}
    for line, rid in line_rows.items():
        if line == "multiple":
            impact_cells["multiple:"] = {"row": rid, "cols": forecast_fy}
            continue
        for y, pid in enumerate(forecast_fy, start=1):
            impact_cells[f"{line}:{y}"] = {"row": rid, "cols": [pid]}

    return {
        "source": base.get("source"),
        "file": base.get("file"),
        "imported_at": base.get("imported_at"),
        "columns": columns,
        "sections": sections,
        "has_consensus": bool(cons and cons.get("lines")),
        "market": {"price": price, "market_cap": mcap, "net_cash": net_cash if bal else None, "ev": ev,
                   "target_pe": target_pe},
        "impact_cells": impact_cells,
        "editable": [p["id"] for p in periods if not p["actual"]
                     and not (p["kind"] == "FY" and len([q for q in qs if q["fy"] == p["fy"]]) == 4)],
    }


def _sub(V, a, b, pid):
    x, y = V.get(a, {}).get(pid), V.get(b, {}).get(pid)
    return x - y if x is not None and y is not None else None


def _add(V, a, b, pid):
    x, y = V.get(a, {}).get(pid), V.get(b, {}).get(pid)
    return x + y if x is not None and y is not None else None
