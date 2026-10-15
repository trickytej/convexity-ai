"""Export a security's model as a ConvexityAI-template workbook.

Layout matches finmodel.parse_convexity, so a downloaded model can be edited in
Excel and uploaded back: a 'Model' sheet (the user's numbers — base model plus
their overrides) and, when available, a 'Consensus Model' sheet.
"""

from __future__ import annotations

import io

from .finmodel import run_engine

_NAVY = "1F3864"
_PCT = "0.0%"
_NUM = "#,##0.0;(#,##0.0)"
_EPS = "0.00;(0.00)"

# (label, kind, line-or-ratio, number format)
_PNL_ROWS = [
    ("Gross margin, %", "ratio", ("gross_profit", "revenue"), _PCT),
    ("Cost of sales", "diff", ("revenue", "gross_profit"), _NUM),
    ("Gross profit", "line", "gross_profit", _NUM),
    ("SG&A, % of revenue", "ratio", ("sga", "revenue"), _PCT),
    ("Selling, general & administrative", "line", "sga", _NUM),
    ("R&D, % of revenue", "ratio", ("rd", "revenue"), _PCT),
    ("Research & development", "line", "rd", _NUM),
    ("EBIT", "line", "operating_income", _NUM),
    ("EBIT margin, %", "ratio", ("operating_income", "revenue"), _PCT),
    ("Add back: depreciation & amortization", "line", "da", _NUM),
    ("EBITDA", "sum", ("operating_income", "da"), _NUM),
    ("Pre-tax income (EBT)", "line", "ebt", _NUM),
    ("Effective tax rate, %", "ratio", ("tax", "ebt"), _PCT),
    ("Income tax provision", "line", "tax", _NUM),
    ("Net Income for EPS Calculation", "line", "net_income", _NUM),
    ("Diluted shares outstanding, mm", "line", "shares", _NUM),
    ("Diluted EPS", "line", "eps", _EPS),
]
_CASH_ROWS = [
    ("Operating cash flow", "line", "ocf", _NUM),
    ("Capex", "line", "capex", _NUM),
]


def _value(get, kind, ref, pid):
    if kind == "line":
        return get(ref, pid)
    a, b = get(ref[0], pid), get(ref[1], pid)
    if a is None or b is None:
        return None
    if kind == "ratio":
        return a / b if b else None
    if kind == "diff":
        return a - b
    return a + b  # sum


def _write_sheet(ws, title: str, periods: list[dict], segments: list[tuple[str, dict]], get, actual_ids: set):
    from openpyxl.styles import Alignment, Font, PatternFill

    head_fill = PatternFill("solid", fgColor=_NAVY)
    sub_fill = PatternFill("solid", fgColor="D9E1F2")
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 44
    ws["B2"] = title
    ws["B2"].font = Font(bold=True, size=13, color=_NAVY)

    ids = [p["id"] for p in periods]
    ws.cell(row=4, column=2, value="USD mm").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=4, column=2).fill = head_fill
    for j, pid in enumerate(ids, start=3):
        c = ws.cell(row=4, column=j, value=pid)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), head_fill, Alignment(horizontal="right")
        ws.column_dimensions[c.column_letter].width = 11
    last_actual = max((j for j, pid in enumerate(ids, start=3) if pid in actual_ids), default=None)
    if last_actual:
        ws.cell(row=5, column=last_actual, value="Actual").font = Font(italic=True, color="808080")
        if last_actual - 2 < len(ids):
            ws.cell(row=5, column=last_actual + 1, value="Forecast").font = Font(italic=True, color="808080")

    r = 6

    def section(name: str) -> None:
        nonlocal r
        c = ws.cell(row=r, column=2, value=name)
        c.font, c.fill = Font(bold=True, color=_NAVY), sub_fill
        r += 1

    def write_row(label: str, fn, fmt: str, bold: bool = False, red: bool = False) -> None:
        nonlocal r
        lc = ws.cell(row=r, column=2, value=label)
        lc.font = Font(bold=bold, color="C00000" if red else "000000")
        for j, pid in enumerate(ids, start=3):
            v = fn(pid)
            if v is not None:
                c = ws.cell(row=r, column=j, value=round(v, 6))
                c.number_format = fmt
                c.font = Font(bold=bold, color="C00000" if red else ("000000" if pid in actual_ids else "1F4E79"))
        r += 1

    def growth(series_get):
        def fn(pid):
            p = next(x for x in periods if x["id"] == pid)
            prev = f"FY{p['fy'] - 1}" if p["kind"] == "FY" else f"Q{p['q']}-{p['fy'] - 1}"
            cur, old = series_get(pid), series_get(prev)
            return cur / old - 1 if cur is not None and old not in (None, 0) else None
        return fn

    section("REVENUE BY SEGMENT")
    for name, sget in segments:
        write_row(f"{name} revenue", sget, _NUM, bold=True)
        write_row("Y/Y growth, %", growth(sget), _PCT)
    write_row("Net Revenue", lambda pid: get("revenue", pid), _NUM, bold=True)
    write_row("Y/Y growth, %", growth(lambda pid: get("revenue", pid)), _PCT)
    r += 1
    section("INCOME STATEMENT")
    for label, kind, ref, fmt in _PNL_ROWS:
        write_row(label, lambda pid, k=kind, rf=ref: _value(get, k, rf, pid), fmt,
                  bold=label in ("Gross profit", "EBIT", "Net Income for EPS Calculation", "Diluted EPS"),
                  red=fmt == _PCT and "margin" in label.lower())
    if any(get("ocf", p["id"]) is not None or get("capex", p["id"]) is not None for p in periods):
        r += 1
        section("CASH FLOW")
        for label, kind, ref, fmt in _CASH_ROWS:
            write_row(label, lambda pid, k=kind, rf=ref: _value(get, k, rf, pid), fmt)
    ws.freeze_panes = "C5"


def export_workbook(base: dict, overrides: dict, company: str) -> bytes:
    """The user's model (+ consensus if present) as .xlsx bytes."""
    from openpyxl import Workbook

    V, _, _ = run_engine(base, overrides)
    periods = sorted(base["periods"], key=lambda p: (p["fy"], p.get("q", 5)))
    actual_ids = {p["id"] for p in periods if p["actual"]}

    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    _write_sheet(
        ws, f"{company} — ConvexityAI model", periods,
        [(s["name"], lambda pid, k=f"seg:{s['id']}": V.get(k, {}).get(pid)) for s in base["segments"]],
        lambda line, pid: V.get(line, {}).get(pid), actual_ids,
    )
    cons = base.get("consensus")
    if cons and cons.get("lines"):
        CL = cons["lines"]
        cons_ids = sorted({pid for ser in CL.values() for pid in ser},
                          key=lambda pid: next(((p["fy"], p.get("q", 5)) for p in periods if p["id"] == pid), (9999, 9)))
        cperiods = [p for p in periods if p["id"] in cons_ids]
        _write_sheet(
            wb.create_sheet("Consensus Model"), f"{company} — consensus", cperiods,
            [(name, lambda pid, ser=ser: ser.get(pid)) for name, ser in (cons.get("segments") or {}).items()],
            lambda line, pid: CL.get(line, {}).get(pid), actual_ids,
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
